"""Render the package's markdown as Word and PDF.

A report nobody can open is not a report. This machine has no pandoc and no
LibreOffice, so the conversion is done here: markdown is parsed once into a
small block model, and two renderers walk it -- `python-docx` for `.docx` and
`reportlab` for `.pdf`. Both are pure Python, so the export cannot break
because a system tool was upgraded or is missing.

The parser is deliberately small: headings, paragraphs, bullet and numbered
lists, tables, fenced code, block quotes and rules, with bold, italic and
inline code inside them. That is the whole of what the package writes. Anything
it does not recognise is emitted as plain text rather than dropped, because a
filing document losing a line silently is worse than one that looks plain.
"""

from __future__ import annotations

import html
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Final, Literal

log = logging.getLogger("ipfe")

__all__ = ["Block", "ExportError", "export_markdown", "parse_markdown", "to_docx", "to_pdf"]

BlockKind = Literal["heading", "para", "bullet", "number", "table", "code", "rule", "quote"]

_HEADING_RE: Final[re.Pattern[str]] = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE: Final[re.Pattern[str]] = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBER_RE: Final[re.Pattern[str]] = re.compile(r"^\s*(\d+)[.)]\s+(.*)$")
_QUOTE_RE: Final[re.Pattern[str]] = re.compile(r"^>\s?(.*)$")
_RULE_RE: Final[re.Pattern[str]] = re.compile(r"^\s*([-*_])\1{2,}\s*$")
_FENCE_RE: Final[re.Pattern[str]] = re.compile(r"^\s*(```|~~~)")
_TABLE_SEP_RE: Final[re.Pattern[str]] = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")
# Inline: **bold**, *italic* / _italic_, `code`. Links keep their text.
_INLINE_RE: Final[re.Pattern[str]] = re.compile(
    r"(\*\*.+?\*\*|__.+?__|\*[^*\n]+?\*|_[^_\n]+?_|`[^`\n]+?`|\[[^\]]+\]\([^)]*\))"
)


class ExportError(RuntimeError):
    """A document could not be rendered. Never fatal to a run: the markdown stands."""


@dataclass
class Block:
    """One renderable unit of the document."""

    kind: BlockKind
    text: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)


@dataclass
class Run:
    """A span of inline text with its styling."""

    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False


# ---------------------------------------------------------------------------
# parsing
# ---------------------------------------------------------------------------


def inline_runs(text: str) -> list[Run]:
    """Split one line into styled runs. Unmatched markers stay literal."""
    runs: list[Run] = []
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if (
            part.startswith("**")
            and part.endswith("**")
            and len(part) > 4
            or part.startswith("__")
            and part.endswith("__")
            and len(part) > 4
        ):
            runs.append(Run(part[2:-2], bold=True))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append(Run(part[1:-1], code=True))
        elif part.startswith("[") and "](" in part:
            runs.append(Run(part[1 : part.index("](")]))
        elif (
            part.startswith("*")
            and part.endswith("*")
            and len(part) > 2
            or part.startswith("_")
            and part.endswith("_")
            and len(part) > 2
        ):
            runs.append(Run(part[1:-1], italic=True))
        else:
            runs.append(Run(part))
    return runs or [Run(text)]


def _table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_markdown(text: str) -> list[Block]:
    """Markdown to blocks. Unknown syntax survives as a paragraph."""
    blocks: list[Block] = []
    lines = (text or "").splitlines()
    index = 0
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block("para", " ".join(paragraph).strip()))
            paragraph.clear()

    while index < len(lines):
        line = lines[index]

        if _FENCE_RE.match(line):
            flush()
            fence = _FENCE_RE.match(line).group(1)  # type: ignore[union-attr]
            index += 1
            body: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith(fence):
                body.append(lines[index])
                index += 1
            blocks.append(Block("code", "\n".join(body)))
            index += 1
            continue

        if not line.strip():
            flush()
            index += 1
            continue

        if _RULE_RE.match(line):
            flush()
            blocks.append(Block("rule"))
            index += 1
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            flush()
            blocks.append(Block("heading", heading.group(2).strip(), level=len(heading.group(1))))
            index += 1
            continue

        # A table is a header row, a separator, then body rows.
        if (
            line.lstrip().startswith("|")
            and index + 1 < len(lines)
            and _TABLE_SEP_RE.match(lines[index + 1])
        ):
            flush()
            rows = [_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append(_table_row(lines[index]))
                index += 1
            blocks.append(Block("table", rows=rows))
            continue

        quote = _QUOTE_RE.match(line)
        if quote:
            flush()
            blocks.append(Block("quote", quote.group(1).strip()))
            index += 1
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            flush()
            indent = len(line) - len(line.lstrip())
            blocks.append(Block("bullet", bullet.group(1).strip(), level=min(indent // 2, 3)))
            index += 1
            continue

        numbered = _NUMBER_RE.match(line)
        if numbered:
            flush()
            blocks.append(Block("number", numbered.group(2).strip()))
            index += 1
            continue

        paragraph.append(line.strip())
        index += 1

    flush()
    return blocks


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def to_docx(blocks: list[Block], target: Path, title: str = "") -> Path:
    """Write the blocks as a Word document."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ExportError(f"python-docx is not available: {exc}") from exc

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    if title:
        document.add_heading(title, level=0)

    for block in blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=min(block.level, 4))
        elif block.kind == "rule":
            document.add_paragraph()
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.font.name = "Menlo"
            run.font.size = Pt(9)
        elif block.kind == "table" and block.rows:
            columns = max(len(row) for row in block.rows)
            table = document.add_table(rows=0, cols=columns)
            table.style = "Table Grid"
            for index, row in enumerate(block.rows):
                cells = table.add_row().cells
                for column in range(columns):
                    text = row[column] if column < len(row) else ""
                    cells[column].text = ""
                    _write_runs(cells[column].paragraphs[0], text, bold=index == 0)
        else:
            style_name = {
                "bullet": "List Bullet",
                "number": "List Number",
                "quote": "Intense Quote",
            }.get(block.kind)
            paragraph = (
                document.add_paragraph(style=style_name) if style_name else document.add_paragraph()
            )
            _write_runs(paragraph, block.text)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))
    return target


def _write_runs(paragraph: Any, text: str, bold: bool = False) -> None:
    from docx.shared import Pt

    for run in inline_runs(text):
        added = paragraph.add_run(run.text)
        added.bold = run.bold or bold
        added.italic = run.italic
        if run.code:
            added.font.name = "Menlo"
            added.font.size = Pt(9)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _markup(text: str) -> str:
    """Inline runs as reportlab markup, with everything else escaped."""
    out: list[str] = []
    for run in inline_runs(text):
        body = html.escape(run.text)
        if run.code:
            body = f'<font face="Courier">{body}</font>'
        if run.bold:
            body = f"<b>{body}</b>"
        if run.italic:
            body = f"<i>{body}</i>"
        out.append(body)
    return "".join(out) or "&nbsp;"


def to_pdf(blocks: list[Block], target: Path, title: str = "") -> Path:
    """Write the blocks as a PDF."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            HRFlowable,
            ListFlowable,
            ListItem,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ExportError(f"reportlab is not available: {exc}") from exc

    sheet = getSampleStyleSheet()
    body = ParagraphStyle(
        "IpfeBody", parent=sheet["BodyText"], fontSize=9.5, leading=13, alignment=TA_LEFT
    )
    code = ParagraphStyle(
        "IpfeCode", parent=body, fontName="Courier", fontSize=8, leading=10, leftIndent=8
    )
    quote = ParagraphStyle(
        "IpfeQuote", parent=body, leftIndent=12, textColor=colors.HexColor("#444")
    )
    headings = {
        level: ParagraphStyle(
            f"IpfeH{level}",
            parent=sheet["Heading1"],
            fontSize=max(9.5, 19 - 2 * level),
            leading=max(12, 22 - 2 * level),
            spaceBefore=10 if level > 1 else 14,
            spaceAfter=4,
        )
        for level in range(1, 7)
    }

    target.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(target),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title or target.stem,
    )
    width = document.width
    story: list[Any] = []
    if title:
        story += [Paragraph(html.escape(title), headings[1]), Spacer(1, 6)]

    pending: list[Any] = []

    def flush_list() -> None:
        if pending:
            story.append(ListFlowable(list(pending), bulletType="bullet", leftIndent=14))
            story.append(Spacer(1, 4))
            pending.clear()

    for block in blocks:
        if block.kind not in {"bullet", "number"}:
            flush_list()
        if block.kind == "heading":
            story.append(Paragraph(_markup(block.text), headings[min(block.level, 6)]))
        elif block.kind == "para":
            story.append(Paragraph(_markup(block.text), body))
            story.append(Spacer(1, 4))
        elif block.kind in {"bullet", "number"}:
            pending.append(ListItem(Paragraph(_markup(block.text), body), leftIndent=14))
        elif block.kind == "quote":
            story.append(Paragraph(_markup(block.text), quote))
            story.append(Spacer(1, 4))
        elif block.kind == "code":
            for line in block.text.splitlines() or [""]:
                story.append(Paragraph(html.escape(line) or "&nbsp;", code))
            story.append(Spacer(1, 6))
        elif block.kind == "rule":
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", color=colors.HexColor("#bbbbbb")))
            story.append(Spacer(1, 6))
        elif block.kind == "table" and block.rows:
            columns = max(len(row) for row in block.rows)
            data = [
                [Paragraph(_markup(row[i] if i < len(row) else ""), body) for i in range(columns)]
                for row in block.rows
            ]
            table = Table(data, colWidths=[width / columns] * columns, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 8))
    flush_list()
    if not story:
        story.append(Paragraph("(empty document)", body))
    # Long claim sets and findings reports run to hundreds of pages; splitting
    # tables across pages is handled by platypus, so nothing else is needed.
    del PageBreak
    document.build(story)
    return target


# ---------------------------------------------------------------------------
# entry point
# ---------------------------------------------------------------------------


def export_markdown(
    source: Path, out_dir: Path, title: str = "", stem: str | None = None
) -> list[Path]:
    """Write `<stem>.docx` and `<stem>.pdf` beside each other. Returns what was written.

    A failure in one format never prevents the other: the caller gets whatever
    could be rendered, and the markdown original is always on disk regardless.
    """
    blocks = parse_markdown(source.read_text(encoding="utf-8"))
    name = stem or source.stem
    written: list[Path] = []
    for renderer, suffix in ((to_docx, ".docx"), (to_pdf, ".pdf")):
        try:
            written.append(renderer(blocks, out_dir / f"{name}{suffix}", title))
        except Exception as exc:  # deliberately broad: an export is never load-bearing
            log.warning("could not render %s as %s: %s", source.name, suffix, exc)
    return written
