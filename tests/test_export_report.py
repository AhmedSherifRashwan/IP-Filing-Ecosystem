"""A report nobody can open is not a report: markdown, Word and PDF."""

from __future__ import annotations

from pathlib import Path

import pytest

from ecosystem.contracts import Finding
from ecosystem.dossier import Dossier
from ecosystem.export import Block, export_markdown, inline_runs, parse_markdown, to_docx, to_pdf
from ecosystem.report import findings_report_markdown
from ecosystem.state import MatterState, new_state

SAMPLE = """# Title

Some **bold** and *italic* and `code` text.

## Table

| id | severity | note |
|---|---|---|
| F-1 | critical | it breaks |

- first bullet
- second bullet

1. first step
2. second step

> a quotation

```
literal code
```

---

Closing paragraph.
"""


def _finding(**over: object) -> Finding:
    data = {
        "id": "F-32-1",
        "source_prompt": "32",
        "category": "novelty",
        "severity": "critical",
        "title": "Claim 1 is anticipated",
        "claim_refs": ["1"],
        "spec_refs": ["[0042]"],
        "attack": "The admitted two-rate core reads on every limitation.",
        "least_narrowing_action": "Recite the purge gap positively.",
        "requires_new_disclosure": True,
        "suggested_target": "P5",
        "reviewers": ["ip-reviewer"],
        "sources": ["28", "32"],
        "consensus": False,
    }
    data.update(over)
    return Finding.model_validate(data)


# ---------------------------------------------------------------------------
# markdown parsing
# ---------------------------------------------------------------------------


def test_every_block_kind_survives_parsing() -> None:
    blocks = parse_markdown(SAMPLE)
    kinds = [block.kind for block in blocks]
    for kind in ("heading", "para", "table", "bullet", "number", "quote", "code", "rule"):
        assert kind in kinds, kind
    table = next(block for block in blocks if block.kind == "table")
    assert table.rows[0] == ["id", "severity", "note"]
    assert table.rows[1] == ["F-1", "critical", "it breaks"]


def test_inline_styles_are_recognised_and_unknown_markers_stay_literal() -> None:
    runs = inline_runs("plain **bold** *ital* `code` 5 * 3 = 15")
    assert [(r.text, r.bold, r.italic, r.code) for r in runs][:4] == [
        ("plain ", False, False, False),
        ("bold", True, False, False),
        (" ", False, False, False),
        ("ital", False, True, False),
    ]
    assert "5 * 3 = 15" in "".join(r.text for r in runs)


def test_an_unparsed_line_is_kept_as_text_not_dropped() -> None:
    blocks = parse_markdown("<!-- a comment -->\n\nreal text")
    assert any("comment" in block.text for block in blocks)
    assert any("real text" in block.text for block in blocks)


# ---------------------------------------------------------------------------
# rendering
# ---------------------------------------------------------------------------


def test_word_and_pdf_are_written_and_are_not_empty(tmp_path: Path) -> None:
    source = tmp_path / "doc.md"
    source.write_text(SAMPLE, encoding="utf-8")
    written = export_markdown(source, tmp_path / "readable", title="Matter — Doc")

    assert {path.suffix for path in written} == {".docx", ".pdf"}
    for path in written:
        assert path.stat().st_size > 2000, path

    from docx import Document

    document = Document(str(tmp_path / "readable" / "doc.docx"))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Closing paragraph." in text
    assert document.tables and document.tables[0].rows[1].cells[0].text == "F-1"
    assert (tmp_path / "readable" / "doc.pdf").read_bytes().startswith(b"%PDF")


def test_a_renderer_failure_still_leaves_the_other_format(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "doc.md"
    source.write_text("# Hi\n\nbody\n", encoding="utf-8")

    def boom(*_args: object, **_kw: object) -> Path:
        raise RuntimeError("no docx today")

    monkeypatch.setattr("ecosystem.export.to_docx", boom)
    written = export_markdown(source, tmp_path / "out")
    assert [path.suffix for path in written] == [".pdf"]


def test_pathological_content_does_not_break_a_renderer(tmp_path: Path) -> None:
    ugly = "# T\n\n" + "| a | b |\n|---|---|\n" + "| <x> & 'q' | **unclosed |\n" * 40
    blocks = parse_markdown(ugly)
    assert to_docx(blocks, tmp_path / "u.docx").stat().st_size > 0
    assert to_pdf(blocks, tmp_path / "u.pdf").read_bytes().startswith(b"%PDF")
    assert to_pdf([Block("para", "")], tmp_path / "empty.pdf").exists()


# ---------------------------------------------------------------------------
# the report
# ---------------------------------------------------------------------------


@pytest.fixture
def matter(tmp_path: Path, settings) -> tuple[Dossier, MatterState]:
    root = settings.dossier_root("reported")
    dossier = Dossier(root, settings).create("reported", "A valve", ["US"])
    return dossier, new_state("reported", "A valve", ["US"])


def test_the_report_states_the_numbers_and_every_severe_finding(matter) -> None:
    dossier, state = matter
    findings = [
        _finding(),
        _finding(id="F-33-1", severity="high", category="design-around", consensus=True),
        _finding(id="F-38-9", severity="low", category="formal", title="Numeral 214 drifts"),
    ]
    text = findings_report_markdown(dossier, state, findings)

    assert "3 distinct findings" in text
    assert "## Critical findings (1)" in text
    assert "## High findings (1)" in text
    assert "F-32-1 — Claim 1 is anticipated" in text
    assert "review steps 28, 32" in text  # merged restatements are shown
    assert "both reviewers agreed" in text
    assert "Numeral 214 drifts" in text  # low findings appear in the table
    assert "No prior art was supplied" in text  # limits are stated
    assert "may authorise a filing" in text


def test_the_report_is_honest_when_there_is_nothing_severe(matter) -> None:
    dossier, state = matter
    text = findings_report_markdown(dossier, state, [_finding(severity="medium")])
    assert "## Critical findings (0)" in text
    assert "No finding was rated critical." in text


def test_a_huge_table_is_flowed_rather_than_lost(tmp_path: Path) -> None:
    """reportlab cannot split a row taller than the page; nothing may be dropped."""
    rows = "\n".join(f"| F-{i} | critical | {'word ' * 200} |" for i in range(60))
    markdown = f"# Residual\n\n| id | severity | note |\n|---|---|---|\n{rows}\n"
    blocks = parse_markdown(markdown)
    pdf = to_pdf(blocks, tmp_path / "big.pdf")
    assert pdf.read_bytes().startswith(b"%PDF")
    assert pdf.stat().st_size > 10_000
    # Word keeps it as a real table: Word paginates rows by itself.
    from docx import Document

    to_docx(blocks, tmp_path / "big.docx")
    assert Document(str(tmp_path / "big.docx")).tables[0].rows.__len__() == 61


def test_a_document_with_its_own_title_is_not_given_a_second_one(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "d.md"
    source.write_text("# The Real Title\n\nbody\n", encoding="utf-8")
    export_markdown(source, tmp_path / "out", title="The Real Title")
    headings = [
        p.text for p in Document(str(tmp_path / "out" / "d.docx")).paragraphs if p.text.strip()
    ]
    assert headings.count("The Real Title") == 1
